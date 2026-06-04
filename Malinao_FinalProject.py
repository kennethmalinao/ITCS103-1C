import tkinter as tk
import openpyxl as op
from tkinter import ttk, messagebox

def display():
    workbook = op.load_workbook("Malinao_Database.xlsx")
    sheet = workbook.active

    for row in table.get_children():
        table.delete(row)

    for row in sheet.iter_rows(min_row=2, values_only=True):
        table.insert("", tk.END, values=row)

def input_validation():
    if guestname_entry.get() == "" or checkin_time_entry.get() == "" or checkout_time_entry.get() == "" or daysreserved_entry.get() == "":
        messagebox.showerror("Error", "All fields are required!")
        return False
    if room_num.get() == "Select Room" or room_type.get() == "Room Type" or checkin_day.get() == "Day" or checkout_day.get() == "Day":
        messagebox.showerror("Error", "Please select correct options from dropdowns!")
        return False
    return True

def compute_days(*args):
    c_day = checkin_day.get()
    co_day = checkout_day.get()

    if c_day == "Day" or co_day == "Day":
        return

    days_map = {"Mon": 1, "Tue": 2, "Wed": 3, "Thu": 4, "Fri": 5, "Sat": 6, "Sun": 7}

    if c_day not in days_map or co_day not in days_map:
        return

    in_num = days_map[c_day]
    out_num = days_map[co_day]

    total_days = out_num - in_num

    if total_days <= 0:
        total_days = total_days + 7

    daysreserved_entry.delete(0, tk.END)
    daysreserved_entry.insert(0, str(total_days))

def saving():
    if not input_validation():
        return

    workbook = op.load_workbook("Malinao_Database.xlsx")
    sheet = workbook.active

    guest = guestname_entry.get()
    roomnum = room_num.get()
    roomtype = room_type.get()
    c_day = checkin_day.get()
    c_time = checkin_time_entry.get()    
    c_ampm = checkin_ampm.get()
    co_day = checkout_day.get()
    co_time = checkout_time_entry.get()  
    co_ampm = checkout_ampm.get()
    downpayment = downpayment_var.get()
    days = daysreserved_entry.get()

    checkin_full = f"{c_day} {c_time} {c_ampm}"
    checkout_full = f"{co_day} {co_time} {co_ampm}"

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if str(row[2]) == str(roomnum) and str(row[4]).startswith(c_day):
            messagebox.showerror("Reserved!", f"Room {roomnum} is already reserved on {c_day}!")
            return 

    new_id = sheet.max_row
    sheet.append([new_id, guest, roomnum, roomtype, checkin_full, checkout_full, downpayment, days])
    workbook.save("Malinao_Database.xlsx")

    messagebox.showinfo("Success", "Record added successfully!")
    display()

def auto_populate(event):
    selected = table.focus()
    values = table.item(selected, "values")

    if values:
        guestname_entry.delete(0, tk.END)
        checkin_time_entry.delete(0, tk.END)
        checkout_time_entry.delete(0, tk.END)
        daysreserved_entry.delete(0, tk.END)

        guestname_entry.insert(0, values[1])
        room_num.set(values[2])
        room_type.set(values[3])

        ci = str(values[4]).split()
        if len(ci) >= 3:
            checkin_day.set(ci[0])
            checkin_time_entry.insert(0, ci[1])
            checkin_ampm.set(ci[2])
        else:
            checkin_time_entry.insert(0, values[4])

        co = str(values[5]).split()
        if len(co) >= 3:
            checkout_day.set(co[0])
            checkout_time_entry.insert(0, co[1])
            checkout_ampm.set(co[2])
        else:
            checkout_time_entry.insert(0, values[5])
        
        downpayment_var.set(values[6])
        daysreserved_entry.insert(0, values[7])

def update():
    selected = table.focus()

    if not selected:
        messagebox.showerror("Error", "Select a record first.")
        return
    
    if not input_validation():
        return
    
    values = table.item(selected, "values")
    record_id = values[0]

    guest = guestname_entry.get()
    roomnum = room_num.get()
    roomtype = room_type.get()
    c_day = checkin_day.get()
    c_time = checkin_time_entry.get()
    c_ampm = checkin_ampm.get()
    co_day = checkout_day.get()
    co_time = checkout_time_entry.get()
    co_ampm = checkout_ampm.get()
    downpayment = downpayment_var.get()
    days = daysreserved_entry.get()

    checkin_full = f"{c_day} {c_time} {c_ampm}"
    checkout_full = f"{co_day} {co_time} {co_ampm}"

    workbook = op.load_workbook("Malinao_Database.xlsx")
    sheet = workbook.active

    for row in sheet.iter_rows(min_row=2):
        if str(row[0].value) != str(record_id):
            if str(row[2].value) == str(roomnum) and str(row[4].value).startswith(c_day):
                messagebox.showerror("Reserved!", f"Room {roomnum} is already reserved on {c_day}!")
                return

    for rows in sheet.iter_rows(min_row=2):
        if str(rows[0].value) == str(record_id):
            rows[1].value = guest
            rows[2].value = roomnum
            rows[3].value = roomtype
            rows[4].value = checkin_full
            rows[5].value = checkout_full
            rows[6].value = downpayment
            rows[7].value = days
            break
    
    workbook.save("Malinao_Database.xlsx")
    messagebox.showinfo("Success", "Record updated successfully!")
    display()

def delete():
    selected = table.focus()

    if not selected:
        messagebox.showerror("Error", "Select a record first!")
        return

    values = table.item(selected, "values")
    record_id = values[0]

    confirm = messagebox.askyesno("Confirm", "Are you sure you want to delete this record?")
    if not confirm:
        return

    workbook = op.load_workbook("Malinao_Database.xlsx")
    sheet = workbook.active

    for i, row in enumerate(sheet.iter_rows(min_row=2), start=2):
        if str(row[0].value) == str(record_id):
            sheet.delete_rows(i)
            break

    workbook.save("Malinao_Database.xlsx")
    messagebox.showinfo("Success", "Record deleted successfully")
    display()

window = tk.Tk()
window.title("Room Reservation System")
window.configure(bg="pink")

downpayment_var = tk.StringVar()
downpayment_var.set("0")

room_type = tk.StringVar()
room_type.set("Room Type")

room_num = tk.StringVar()
room_num.set("Select Room")

checkin_day = tk.StringVar()
checkin_day.set("Day")
checkin_day.trace("w", compute_days)

checkin_ampm = tk.StringVar()
checkin_ampm.set("AM/PM")

checkout_day = tk.StringVar()
checkout_day.set("Day")
checkout_day.trace("w", compute_days)

checkout_ampm = tk.StringVar()
checkout_ampm.set("AM/PM")

title = tk.Label(window, text="Room Reservation System", font=("Arial", 25, "italic"), bg="pink", fg="black")
title.grid(row=0, column=0, columnspan=6, pady=10)

genframe = tk.Frame(window, bg="brown", bd=20, relief="ridge")
genframe.grid(row=1, column=0, columnspan=6, padx=10, pady=10)

guestname_entry = tk.Entry(genframe, font=("Poppins", 12), width=18)
guestname_entry.grid(row=2, column=1, padx=5, pady=(10, 0))
guestname_label = tk.Label(genframe, text="Guest Name", font=("Poppins", 10, "italic"), bg="brown", fg="white")
guestname_label.grid(row=3, column=1)

roomnumber_menu = tk.OptionMenu(genframe, room_num, *range(1, 31))
roomnumber_menu.config(width=12)
roomnumber_menu.grid(row=2, column=2, padx=5, pady=(10, 0))
roomnumber_label = tk.Label(genframe, text="Room Number", font=("Poppins", 10, "italic"), bg="brown", fg="white")
roomnumber_label.grid(row=3, column=2)

roomtype_menu = tk.OptionMenu(genframe, room_type, "Single", "Double", "Deluxe")
roomtype_menu.config(width=12)
roomtype_menu.grid(row=2, column=3, padx=5, pady=(10, 0))
roomtype_label = tk.Label(genframe, text="Room Type", font=("Poppins", 10, "italic"), bg="brown", fg="white")
roomtype_label.grid(row=3, column=3)

checkin_day_menu = tk.OptionMenu(genframe, checkin_day, "Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun")
checkin_day_menu.config(width=5)
checkin_day_menu.grid(row=4, column=1, padx=2, pady=(10, 0))

checkin_time_entry = tk.Entry(genframe, font=("Poppins", 12), width=6)
checkin_time_entry.grid(row=4, column=2, padx=2, pady=(10, 0))

checkin_ampm_menu = tk.OptionMenu(genframe, checkin_ampm, "AM", "PM")
checkin_ampm_menu.config(width=5)
checkin_ampm_menu.grid(row=4, column=3, padx=2, pady=(10, 0))

checkin_label = tk.Label(genframe, text="Check-In (Day / Time / AM-PM)", font=("Poppins", 10, "italic"), bg="brown", fg="white")
checkin_label.grid(row=5, column=1, columnspan=3)

checkout_day_menu = tk.OptionMenu(genframe, checkout_day, "Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun")
checkout_day_menu.config(width=5)
checkout_day_menu.grid(row=6, column=1, padx=2, pady=(10, 0))

checkout_time_entry = tk.Entry(genframe, font=("Poppins", 12), width=6)
checkout_time_entry.grid(row=6, column=2, padx=2, pady=(10, 0))

checkout_ampm_menu = tk.OptionMenu(genframe, checkout_ampm, "AM", "PM")
checkout_ampm_menu.config(width=5)
checkout_ampm_menu.grid(row=6, column=3, padx=2, pady=(10, 0))

checkout_label = tk.Label(genframe, text="Check-Out (Day / Time / AM-PM)", font=("Poppins", 10, "italic"), bg="brown", fg="white")
checkout_label.grid(row=7, column=1, columnspan=3)

down_payment = tk.OptionMenu(genframe, downpayment_var, "0", "1000", "2000", "3000", "Fully Paid")
down_payment.config(width=12)
down_payment.grid(row=8, column=1, padx=5, pady=(10, 0))

down_label = tk.Label(genframe, text="Down Payment", font=("Poppins", 10, "italic"), bg="brown", fg="white")
down_label.grid(row=9, column=1)

daysreserved_entry = tk.Entry(genframe, font=("Poppins", 12), width=10)
daysreserved_entry.grid(row=8, column=2, padx=5, pady=(10, 0))

days_label = tk.Label(genframe, text="Days Reserved", font=("Poppins", 10, "italic"), bg="brown", fg="white")
days_label.grid(row=9, column=2)

button = tk.Button(window, text="Submit", font=("Poppins", 11, "bold"), bg="lightgreen", command=saving)
button.grid(row=10, column=1, pady=10)

update_btn = tk.Button(window, text="Update", command=update)
update_btn.grid(row=10, column=2, pady=10)

delete_btn = tk.Button(window, text="Delete", bg="red", fg="white", command=delete)
delete_btn.grid(row=10, column=3, pady=10)

table_style = ttk.Style()
table_style.configure("Treeview.Heading", font=("Poppins", 10, "bold"), foreground="darkgreen")
table_style.configure("Treeview", font=("Poppins", 10), rowheight=25)

table = ttk.Treeview(window, columns=("ID", "Guest Name", "Room Number", "Room Type", "Check-In", "Check-Out", "Down Payment", "Days Reserved"), show="headings")
for headings in ("ID", "Guest Name", "Room Number", "Room Type", "Check-In", "Check-Out", "Down Payment", "Days Reserved"):
    table.heading(headings, text=headings)
    table.column(headings, width=110, anchor="center")
table.grid(row=11, column=0, columnspan=6, padx=10, pady=10)

table.bind("<<TreeviewSelect>>", auto_populate)

display()
window.mainloop()